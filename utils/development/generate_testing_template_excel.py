import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Read the rules from rules.json
with open('rules.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "Testing Template"

# Define styles
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=11)
rule_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
tester_pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
tester_comment_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Create header row
headers = ['Rule #', 'Rule Name', 'Description']
for i in range(1, 11):
    headers.append(f'Tester {i}\nPass/Fail')
    headers.append(f'Tester {i}\nComments')

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.value = header
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = border

# Set column widths
ws.column_dimensions['A'].width = 8   # Rule #
ws.column_dimensions['B'].width = 35  # Rule Name
ws.column_dimensions['C'].width = 60  # Description

# Tester columns
for i in range(4, 24, 2):  # Columns D, F, H, J, L, N, P, R, T, V (Pass/Fail)
    ws.column_dimensions[get_column_letter(i)].width = 12
for i in range(5, 25, 2):  # Columns E, G, I, K, M, O, Q, S, U, W (Comments)
    ws.column_dimensions[get_column_letter(i)].width = 30

# Add rules
current_row = 2
rule_number = 1

# Create data validation for Pass/Fail dropdown
pass_fail_validation = DataValidation(
    type="list",
    formula1='"Pass,Fail"',
    showErrorMessage=True,
    errorTitle="Invalid Input",
    error="Please select either 'Pass' or 'Fail' from the dropdown"
)
ws.add_data_validation(pass_fail_validation)

for rule in data['rules']:
    rule_name = rule['name']
    
    for case in rule['cases']:
        description = case['description']
        
        # Rule # column
        cell = ws.cell(row=current_row, column=1)
        cell.value = rule_number
        cell.alignment = center_align
        cell.fill = rule_fill
        cell.border = border
        
        # Rule Name column
        cell = ws.cell(row=current_row, column=2)
        cell.value = rule_name
        cell.alignment = left_align
        cell.fill = rule_fill
        cell.border = border
        
        # Description column
        cell = ws.cell(row=current_row, column=3)
        cell.value = description
        cell.alignment = left_align
        cell.border = border
        
        # Tester columns
        for col_num in range(4, 24):
            cell = ws.cell(row=current_row, column=col_num)
            cell.border = border
            cell.alignment = center_align if col_num % 2 == 0 else left_align
            
            # Color code: Pass/Fail columns in green, Comments in yellow
            if col_num % 2 == 0:  # Pass/Fail columns
                cell.fill = tester_pass_fill
                # Add dropdown to Pass/Fail cells
                pass_fail_validation.add(cell)
            else:  # Comments columns
                cell.fill = tester_comment_fill
        
        current_row += 1
    
    rule_number += 1

# Freeze header row
ws.freeze_panes = 'A2'

# Add instructions sheet
instructions_ws = wb.create_sheet("Instructions", 0)
instructions_ws.column_dimensions['A'].width = 100

instructions = [
    "Quranic App Testing Template - Instructions",
    "",
    "Purpose:",
    "This template is designed for testing all Tajweed rules implemented in the Quranic application.",
    "",
    "How to Use:",
    "1. Each row represents one rule case with its description.",
    "2. There are 10 tester columns (Tester 1 through Tester 10).",
    "3. For each tester, there are two columns:",
    "   - Pass/Fail: Select 'Pass' or 'Fail' from the dropdown menu",
    "   - Comments: Add any observations, bugs, or suggestions",
    "",
    "Testing Guidelines:",
    "- Test each rule thoroughly by searching for examples in the Quran",
    "- Verify that the search results match the rule description",
    "- Check if the highlighting/marking is correct",
    "- Test edge cases and boundary conditions",
    "- Note any false positives or false negatives",
    "",
    "Rule Categories Covered:",
    "- Lam Shamsiyyah & Qamariyyah (اللام الشمسية والقمرية)",
    "- Noon & Tanween Rules (أحكام النون الساكنة والتنوين)",
    "- Meem Rules (أحكام الميم الساكنة)",
    "- Qalqalah (قلقلة الحروف)",
    "- Madd Rules (أحكام المدود)",
    "- Waqf Rules (أحكام الوقف)",
    "- Tafkheem & Tarqeeq (التفخيم والترقيق)",
    "",
    "Contact Information:",
    "If you encounter critical bugs or have urgent questions, please contact the development team.",
    "",
    "Thank you for your valuable feedback!"
]

for row_num, instruction in enumerate(instructions, 1):
    cell = instructions_ws.cell(row=row_num, column=1)
    cell.value = instruction
    if row_num == 1:
        cell.font = Font(size=16, bold=True, color="366092")
    elif instruction.endswith(':'):
        cell.font = Font(size=12, bold=True)
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

# Save workbook
output_file = 'Quraan_App_Testing_Template_With_Dropdowns.xlsx'
wb.save(output_file)
print(f"Excel testing template with dropdowns created successfully: {output_file}")
print(f"Total rules to test: {rule_number - 1}")
print(f"Total test cases: {current_row - 2}")